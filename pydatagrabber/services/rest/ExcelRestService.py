from __future__ import annotations
from dataclasses import dataclass, field
from fastapi import FastAPI
from openpyxl import load_workbook

from ...services.ServiceException import ServiceException

from ...utils.FileUtils import FileUtils
from .ExcelRestAPI import ExcelRestAPI
from ...buffers.DictBuffer import DictBuffer
from ..Service import Service
from ..rest.RestService import RestService
from ...grabbers.Grabber import Grabber


@dataclass
class ExcelRestService(RestService):
    """Service for creating a REST API for accessing named Tables in Excel
    """
    
    excel_file : str = field(default=None, metadata={"description" : "path of the excel file to open for tables"})
           
    def __init__(self):
        super().__init__()
        self.named_tables : dict[str, DictBuffer] = dict()
        
    def install(self, grabber : Grabber = None):
        super(Service, self).install()
        self.app = FastAPI(title="DataGrabber ExcelRestService", docs_url="/docs")
        self.__add_cors()
        self.app.include_router(ExcelRestAPI.get_api_router(self))
        self.__init_tables()
        
    def deinstall(self, grabber : Grabber = None):
        super().deinstall(grabber)
        self.app = None
        
    def __init_tables(self):
        if FileUtils.exists_file(self.excel_file):
            # Load the workbook
            wb = load_workbook(self.excel_file, data_only=True)
            for ws in wb.worksheets:
                for table_name, table_range in ws.tables.items():
                    #print("🧾 Table name:", table_name)
                    #print("📐 Range:", table_range)
                    # Access the cells inside the table range
                    cells = ws[table_range]
                    dbuf = DictBuffer()
                    dbuf.id = table_name
                    dbuf.capacity = len(cells)
                    r = 0
                    headers = []
                    for row in cells:
                        #print([cell.value for cell in row]) 
                        if r == 0:
                            headers = [cell.value for cell in row]
                        else:
                            d = dict()
                            h = 0
                            for cell in row:
                                d[headers[h]] = cell.value
                                h = h + 1
                            dbuf.push(d)
                        r = r + 1                       
                         
                    self.named_tables[table_name] = dbuf
        else:
            raise ServiceException("the file " + self.excel_file + " could not be found")