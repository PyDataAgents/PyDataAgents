from dataclasses import dataclass, field

from openpyxl import load_workbook
from openpyxl.worksheet.table import Table

from ..NodeException import NodeException
from ...utils.FileUtils import FileUtils
from ..Action import Action
from ..BufferNode import BufferNode

@dataclass
class ReadExcelTableAction(BufferNode, Action):
    
    excel_file : str = field(default=None, metadata={"description": "path to the excel files to read named table from"})
    table_name : str = field(default=None, metadata={"description": "name of the table inside the excel to read from"})
        
    def execute(self):
        if FileUtils.exists_file(self.excel_file):
            # Load the workbook
            wb = load_workbook(self.excel_file, data_only=True)
            success = False
            for ws in wb.worksheets:
                if self.table_name in ws.tables:
                    t_range : Table = ws.tables[self.table_name]
                    #print("🧾 Table name:", table_name)
                    #print("📐 Range:", table_range)
                    # Access the cells inside the table range                    
                    cells = ws[t_range.ref]
                    r = 0
                    headers = []
                    for row in cells:
                        #print([cell.value for cell in row]) 
                        if r == 0:
                            headers = [cell.value for cell in row]
                        else:
                            d = dict()
                            h = 0
                            for cell in row:
                                d[headers[h]] = cell.value
                                h = h + 1
                            self.buffer.push(d)
                        r = r + 1
                    success = True
                    break
            if not success:
                raise NodeException("No name table '" + self.table_name + "' could be found")
        else:
            raise NodeException("the file " + self.excel_file + " could not be found")