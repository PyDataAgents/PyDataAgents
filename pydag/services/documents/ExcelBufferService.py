from __future__ import annotations
from dataclasses import dataclass, field
from openpyxl import load_workbook
import pandas as pd

from ..ServiceException import ServiceException
from ...utils.FileUtils import FileUtils
from ...buffers.DictBuffer import DictBuffer
from ..Service import Service
from ...agents.Agent import Agent


@dataclass
class ExcelBufferService(Service):
    """
    `Service` for creating `Buffer`s in `Agent` for each named table found in specified Excel file
    """
    
    excel_file : str = field(default=None, metadata={"description" : "path of the excel file to open for tables"})
           
    def __post_init__(self):
        super().__post_init__()
        self._named_tables : list[str] = []
        
    def _on_start(self):
        if self._agent is None:
            raise ServiceException("No " + Agent.__name__ + " was specified, make sure to install the " + self.cname() + " before starting!")
        if FileUtils.exists_file(self.excel_file):
            # Load the workbook
            wb = load_workbook(self.excel_file, data_only=True)
            for ws in wb.worksheets:
                for table_name, table_range in ws.tables.items():
                    #print("🧾 Table name:", table_name)
                    #print("📐 Range:", table_range)
                    # Access the cells inside the table range
                    #print(type(table_range))
                    cells = ws[table_range]
                    dbuf = DictBuffer(id=table_name, capacity=len(cells))
                    self._agent.add_buffer(dbuf)
                    dbuf.install(self._agent)
                    r = 0
                    headers = []
                    for row in cells:
                        #print([cell.value for cell in row]) 
                        if r == 0:
                            headers = [cell.value for cell in row]
                        else:
                            d = dict()
                            h = 0
                            for cell in row:
                                d[headers[h]] = cell.value
                                h = h + 1
                            dbuf.push(d)
                        r = r + 1
                    self._named_tables.append(dbuf.id)
        else:
            raise ServiceException("the file " + self.excel_file + " could not be found")
    
    def _on_stop(self):
        self._named_tables = []
        
    def to_dataframes(self) -> dict[str, pd.DataFrame]:
        """Returns a dictionary of pandas DataFrames for each named table in the excel file

        Returns:
            dict[str, pd.DataFrame]: dictionary of DataFrames
        """
        dfs : dict[str, pd.DataFrame] = dict()
        for name in self._named_tables:
            buf = self._agent.get_buffer(name)
            dfs[name] = pd.DataFrame(buf.data())
        return dfs
    
    def get_named_tables(self) -> list[str]:
        """Returns the list of named tables found in the excel file

        Returns:
            list[str]: list of named tables
        """
        return self._named_tables
